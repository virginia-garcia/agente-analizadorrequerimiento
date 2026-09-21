from datetime import date
from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from modelos import Documento

AZUL = "1F4E79"

_REEMPLAZOS = {"≥": ">=", "≤": "<=", "→": "->", "←": "<-", "≈": "~", "–": "-", "—": "-",
               "“": '"', "”": '"', "‘": "'", "’": "'", "…": "...", "•": "-"}


def _para_pdf(texto: str) -> str:
    """La fuente estándar del PDF solo soporta Latin-1: reemplaza o descarta lo demás."""
    for k, v in _REEMPLAZOS.items():
        texto = texto.replace(k, v)
    return texto.encode("cp1252", errors="ignore").decode("cp1252")


# ---------------------------------------------------------------- Markdown
def a_markdown(doc: Documento, requerimiento: str) -> str:
    L = ["# Documentación funcional\n", f"> **Requerimiento:** {requerimiento}\n",
         f"## Resumen\n{doc.resumen}\n", "## Historias de usuario"]
    for h in doc.historias_de_usuario:
        L += [f"\n### {h.id} · {h.titulo} _(prioridad {h.prioridad})_",
              f"Como **{h.como}**, quiero {h.quiero}, para {h.para}.\n",
              "**Criterios de aceptación:**"] + [f"- {c}" for c in h.criterios_de_aceptacion]
    L += ["\n## Requisitos no funcionales"] + [f"- {r}" for r in doc.requisitos_no_funcionales]
    st = doc.solucion_tecnica
    L += [f"\n## Solución técnica\n{st.enfoque}\n", "**Componentes:**"]
    L += [f"- {c}" for c in st.componentes]
    if st.datos_necesarios:
        L += ["\n**Datos necesarios:**"] + [f"- {d}" for d in st.datos_necesarios]
    L += ["\n## Riesgos"] + [f"- **{r.riesgo}** → {r.mitigacion}" for r in doc.riesgos]
    L += ["\n## Métricas de éxito"] + [f"- {m.metrica}: {m.objetivo}" for m in doc.metricas_de_exito]
    for titulo, items in (("Supuestos", doc.supuestos), ("Fuera de alcance", doc.fuera_de_alcance),
                          ("Preguntas abiertas", doc.preguntas_abiertas)):
        if items:
            L += [f"\n## {titulo}"] + [f"- {i}" for i in items]
    return "\n".join(L) + "\n"


# -------------------------------------------------------------------- Word
def _sombrear(celda, hex_color: str):
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _tabla_word(word, encabezados: list[str], filas: list[list[str]]):
    t = word.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    for i, h in enumerate(encabezados):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _sombrear(c, AZUL)
    for fila in filas:
        celdas = t.add_row().cells
        for i, valor in enumerate(fila):
            celdas[i].text = valor
    word.add_paragraph()


def a_docx(doc: Documento, requerimiento: str) -> bytes:
    w = WordDocument()
    w.styles["Normal"].font.name = "Calibri"
    w.styles["Normal"].font.size = Pt(11)

    titulo = w.add_heading("Documentación funcional", level=0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    w.add_paragraph(f"Generado el {date.today():%d/%m/%Y}").runs[0].italic = True

    w.add_heading("Requerimiento", level=1)
    w.add_paragraph(requerimiento)

    w.add_heading("Resumen", level=1)
    w.add_paragraph(doc.resumen)

    w.add_heading("Historias de usuario", level=1)
    for h in doc.historias_de_usuario:
        w.add_heading(f"{h.id} · {h.titulo} (prioridad {h.prioridad})", level=2)
        p = w.add_paragraph()
        p.add_run("Como ")
        p.add_run(h.como).bold = True
        p.add_run(f", quiero {h.quiero}, para {h.para}.")
        w.add_paragraph().add_run("Criterios de aceptación:").bold = True
        for c in h.criterios_de_aceptacion:
            w.add_paragraph(c, style="List Bullet")

    w.add_heading("Requisitos no funcionales", level=1)
    for r in doc.requisitos_no_funcionales:
        w.add_paragraph(r, style="List Bullet")

    st = doc.solucion_tecnica
    w.add_heading("Solución técnica", level=1)
    w.add_paragraph(st.enfoque)
    if st.componentes:
        w.add_heading("Componentes", level=2)
        for c in st.componentes:
            w.add_paragraph(c, style="List Bullet")
    if st.datos_necesarios:
        w.add_heading("Datos necesarios", level=2)
        for d in st.datos_necesarios:
            w.add_paragraph(d, style="List Bullet")

    w.add_heading("Riesgos", level=1)
    _tabla_word(w, ["Riesgo", "Mitigación"], [[r.riesgo, r.mitigacion] for r in doc.riesgos])

    w.add_heading("Métricas de éxito", level=1)
    _tabla_word(w, ["Métrica", "Objetivo"], [[m.metrica, m.objetivo] for m in doc.metricas_de_exito])

    for titulo, items in (("Supuestos", doc.supuestos), ("Fuera de alcance", doc.fuera_de_alcance),
                          ("Preguntas abiertas", doc.preguntas_abiertas)):
        if items:
            w.add_heading(titulo, level=1)
            for i in items:
                w.add_paragraph(i, style="List Bullet")

    buf = BytesIO()
    w.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------- PDF
def a_pdf(doc: Documento, requerimiento: str) -> bytes:
    base = getSampleStyleSheet()
    color = colors.HexColor(f"#{AZUL}")
    titulo = ParagraphStyle("T", parent=base["Title"], textColor=color, alignment=0, fontSize=22)
    h1 = ParagraphStyle("H1", parent=base["Heading1"], textColor=color, fontSize=14, spaceBefore=14)
    h2 = ParagraphStyle("H2", parent=base["Heading2"], textColor=color, fontSize=12)
    normal = ParagraphStyle("N", parent=base["Normal"], fontSize=10.5, leading=15)
    negrita = ParagraphStyle("B", parent=normal, fontName="Helvetica-Bold")
    celda = ParagraphStyle("C", parent=normal, fontSize=10, leading=13)
    celda_h = ParagraphStyle("CH", parent=celda, textColor=colors.white, fontName="Helvetica-Bold")
    P = lambda t, s=normal: Paragraph(escape(_para_pdf(t)), s)

    def bullets(items):
        return ListFlowable([ListItem(P(i)) for i in items], bulletType="bullet", leftIndent=14)

    def tabla(encabezados, filas):
        data = [[Paragraph(h, celda_h) for h in encabezados]]
        data += [[P(v, celda) for v in f] for f in filas]
        t = Table(data, colWidths=[8.5 * cm, 8.5 * cm], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), color),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#BBBBBB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FA")]),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return t

    s = [P("Documentación funcional", titulo),
         P(f"Generado el {date.today():%d/%m/%Y}", ParagraphStyle("D", parent=normal, textColor=colors.grey)),
         P("Requerimiento", h1), P(requerimiento),
         P("Resumen", h1), P(doc.resumen),
         P("Historias de usuario", h1)]
    for h in doc.historias_de_usuario:
        s += [P(f"{h.id} · {h.titulo} (prioridad {h.prioridad})", h2),
              P(f"Como {h.como}, quiero {h.quiero}, para {h.para}."),
              Spacer(1, 4), P("Criterios de aceptación:", negrita), bullets(h.criterios_de_aceptacion)]
    s += [P("Requisitos no funcionales", h1), bullets(doc.requisitos_no_funcionales),
          P("Solución técnica", h1), P(doc.solucion_tecnica.enfoque)]
    if doc.solucion_tecnica.componentes:
        s += [P("Componentes", h2), bullets(doc.solucion_tecnica.componentes)]
    if doc.solucion_tecnica.datos_necesarios:
        s += [P("Datos necesarios", h2), bullets(doc.solucion_tecnica.datos_necesarios)]
    s += [P("Riesgos", h1), tabla(["Riesgo", "Mitigación"], [[r.riesgo, r.mitigacion] for r in doc.riesgos]),
          P("Métricas de éxito", h1),
          tabla(["Métrica", "Objetivo"], [[m.metrica, m.objetivo] for m in doc.metricas_de_exito])]
    for titulo, items in (("Supuestos", doc.supuestos), ("Fuera de alcance", doc.fuera_de_alcance),
                          ("Preguntas abiertas", doc.preguntas_abiertas)):
        if items:
            s += [P(titulo, h1), bullets(items)]

    def pie(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Página {d.page}")
        canvas.restoreState()

    buf = BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                      topMargin=2 * cm, bottomMargin=2 * cm,
                      title="Documentación funcional").build(s, onFirstPage=pie, onLaterPages=pie)
    return buf.getvalue()
